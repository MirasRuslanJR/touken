"""Разбор списка учеников из файла — .xlsx или .csv.

Делается на сервере: раньше файл парсила библиотека SheetJS с CDN, и в
школьной сети без доступа к cdnjs импорт не работал вовсе. openpyxl уже стоит
ради экспорта, так что новой зависимости это не приносит.

Формат гибкий намеренно: школьные списки приходят в разном виде. Понимаются
заголовки («имя», «фамилия», «пол», «дата рождения») в любом регистре; если
заголовков нет — берётся первый столбец как ФИО.
"""
import csv
import io
import re
from datetime import date, datetime

NAME_HEADERS = ("фио", "имя", "ученик", "ф.и.о.", "name", "full_name", "фамилия")
GENDER_HEADERS = ("пол", "gender", "sex")
BIRTH_HEADERS = ("дата рождения", "дата", "др", "birth", "birth_date", "дата_рождения")

MALE = ("м", "муж", "мужской", "m", "male", "ұл", "ер")
FEMALE = ("ж", "жен", "женский", "f", "female", "қыз", "әйел")


def _norm(value):
    return re.sub(r"\s+", " ", str(value or "")).strip()


def _gender(value):
    v = _norm(value).lower().rstrip(".")
    if v in MALE:
        return "m"
    if v in FEMALE:
        return "f"
    return None


def _birth(value):
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = _norm(value)
    for fmt in ("%Y-%m-%d", "%d.%m.%Y", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(text, fmt).date().isoformat()
        except ValueError:
            continue
    return None


def _header_map(row):
    """Индексы колонок по заголовкам. None, если заголовков нет."""
    lowered = [_norm(c).lower() for c in row]
    idx = {}
    for i, cell in enumerate(lowered):
        if cell in NAME_HEADERS and "name" not in idx:
            idx["name"] = i
        elif cell in GENDER_HEADERS and "gender" not in idx:
            idx["gender"] = i
        elif cell in BIRTH_HEADERS and "birth" not in idx:
            idx["birth"] = i
    return idx if "name" in idx else None


def _rows_to_students(rows):
    rows = [r for r in rows if any(_norm(c) for c in r)]
    if not rows:
        return []

    mapping = _header_map(rows[0])
    if mapping is not None:
        body = rows[1:]
    else:
        # Заголовков нет — первый столбец это ФИО, второй пол, третий дата.
        mapping = {"name": 0, "gender": 1, "birth": 2}
        body = rows

    out = []
    for r in body:
        def cell(key):
            i = mapping.get(key)
            return r[i] if i is not None and i < len(r) else None

        name = _norm(cell("name"))
        if not name or len(name) > 200:
            continue
        out.append({
            "full_name": name,
            "gender": _gender(cell("gender")),
            "birth_date": _birth(cell("birth")),
        })
    return out


def parse_csv(raw):
    # Школьные выгрузки в Казахстане чаще всего в UTF-8 или cp1251.
    for encoding in ("utf-8-sig", "utf-8", "cp1251"):
        try:
            text = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise ValueError("Не удалось определить кодировку файла. Сохраните его в UTF-8.")

    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=";,\t")
    except csv.Error:
        dialect = csv.excel
        dialect.delimiter = ";" if sample.count(";") > sample.count(",") else ","
    return _rows_to_students(list(csv.reader(io.StringIO(text), dialect)))


def parse_xlsx(raw):
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception:
        raise ValueError("Не удалось прочитать файл Excel. Сохраните его в формате .xlsx или .csv.")
    ws = wb.active
    rows = [list(r) for r in ws.iter_rows(max_col=8, max_row=1000, values_only=True)]
    wb.close()
    return _rows_to_students(rows)


def parse_student_file(filename, raw):
    lower = (filename or "").lower()
    if lower.endswith(".csv") or lower.endswith(".txt"):
        return parse_csv(raw)
    if lower.endswith(".xlsx") or lower.endswith(".xlsm"):
        return parse_xlsx(raw)
    if lower.endswith(".xls"):
        raise ValueError("Старый формат .xls не поддерживается. Сохраните файл как .xlsx или .csv.")
    # Расширения нет — пробуем по сигнатуре: xlsx это zip.
    if raw[:2] == b"PK":
        return parse_xlsx(raw)
    return parse_csv(raw)
