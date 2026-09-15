"""Отчёты в Excel — собираются на сервере.

Раньше файл собирался в браузере библиотекой SheetJS с CDN: без интернета
кнопка «Экспорт» просто не работала, а школьные сети cdnjs фильтруют. Теперь
кабинету для экспорта интернет не нужен вовсе.

Отчёт по классу намеренно не содержит авторства выборов: те же правила, что и
в интерфейсе (см. заголовок app/analytics.py). Номинации «часто остаётся
один» — числом и только начиная с порога.
"""
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="2563EB")
HEADER_FONT = Font(color="FFFFFF", bold=True)
TITLE_FONT = Font(bold=True, size=13)

STATUS_TEXT = {"isolate": "изолят", "unknown": "нет данных", "connected": "есть связи"}
RELIABILITY_TEXT = {"high": "высокая", "medium": "средняя", "low": "низкая"}


def _alone_text(metrics, threshold):
    votes = metrics.get("alone_votes") or 0
    if not votes:
        return 0
    return votes if metrics.get("alone_reportable") else "менее %d" % threshold


def _autosize(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def class_report(class_name, survey, analysis, students, threshold=3):
    """Книга Excel по одному срезу: сводка класса + таблица по ученикам."""
    gm = analysis["graph_metrics"]
    per = analysis["per_student"]

    wb = Workbook()
    ws = wb.active
    ws.title = (survey.get("title") or "Срез")[:28].replace("/", " ").replace("\\", " ")

    ws["A1"] = "Отчёт по классу: %s" % class_name
    ws["A1"].font = TITLE_FONT
    ws["A2"] = "Срез: %s от %s" % (survey.get("title", ""), survey.get("conducted_on", ""))

    summary = [
        ("Учеников", gm["students"]),
        ("Прошли опрос", gm["responded"]),
        ("Явка", "%d%%" % round(gm["participation"] * 100)),
        ("Достоверность среза", RELIABILITY_TEXT.get(gm["reliability"], gm["reliability"])),
        ("Изоляты", gm["isolates"]),
        ("Без данных (низкое покрытие)", gm["unknown"]),
        ("Взаимные пары", gm["mutual_pairs"]),
        ("Плотность (среди ответивших)", gm["density"]),
        ("Взаимность", gm["reciprocity"]),
        ("Сплочённость", gm["cohesion"]),
        ("Индекс связности класса", gm["wellbeing_index"] if gm["wellbeing_index"] is not None else "не считается"),
    ]
    row = 4
    for label, value in summary:
        ws.cell(row=row, column=1, value=label)
        ws.cell(row=row, column=2, value=value)
        row += 1

    if gm["reliability"] == "low":
        ws.cell(row=row + 1, column=1,
                value="Внимание: опрос прошли менее 70% класса. Показатели считаются по ответившим, "
                      "статус «изолят» в этом срезе не присваивается.")
        ws.cell(row=row + 1, column=1).alignment = Alignment(wrap_text=True)
        row += 2

    head_row = row + 2
    headers = ["Ученик", "Статус", "Входящие", "Исходящие", "Взаимные",
               "«Часто один»", "Degree centrality", "Betweenness", "Группа", "Прошёл опрос"]
    for col, title in enumerate(headers, start=1):
        cell = ws.cell(row=head_row, column=col, value=title)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT

    ordered = sorted(students, key=lambda s: (per.get(str(s["id"]), {}).get("in_degree", 0), s["full_name"]))
    for i, s in enumerate(ordered, start=1):
        m = per.get(str(s["id"]), {})
        r = head_row + i
        values = [
            s["full_name"],
            STATUS_TEXT.get(m.get("status"), "—"),
            m.get("in_degree", 0),
            m.get("out_degree", 0),
            m.get("mutual", 0),
            _alone_text(m, threshold),
            m.get("degree_centrality", 0),
            m.get("betweenness", 0),
            (m.get("community", 0) or 0) + 1,
            "да" if m.get("responded") else "нет",
        ]
        for col, value in enumerate(values, start=1):
            ws.cell(row=r, column=col, value=value)

    _autosize(ws, [28, 14, 11, 11, 11, 14, 18, 14, 9, 14])
    ws.freeze_panes = ws.cell(row=head_row + 1, column=1)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def school_report(school_name, rows):
    """Книга Excel по школе: строка на класс, без данных отдельных учеников.

    Именно этот файл уходит завучу — поимённой информации в нём нет by design.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Школа"

    ws["A1"] = "Сводка по школе: %s" % school_name
    ws["A1"].font = TITLE_FONT

    headers = ["Класс", "Психолог", "Учеников", "Последний срез", "Явка",
               "Достоверность", "Индекс связности", "Изоляты", "Требуют внимания",
               "Профилактика: план", "Профилактика: проведено"]
    for col, title in enumerate(headers, start=1):
        cell = ws.cell(row=3, column=col, value=title)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT

    for i, r in enumerate(rows, start=1):
        values = [
            r["class_name"], r.get("psychologist") or "—", r["students"],
            r.get("last_survey") or "нет срезов",
            "%d%%" % round((r.get("participation") or 0) * 100) if r.get("last_survey") else "—",
            RELIABILITY_TEXT.get(r.get("reliability"), "—"),
            r["wellbeing_index"] if r.get("wellbeing_index") is not None else "не считается",
            r.get("isolates", 0),
            r.get("open_alerts", 0),
            r.get("activities_planned", 0),
            r.get("activities_done", 0),
        ]
        for col, value in enumerate(values, start=1):
            ws.cell(row=3 + i, column=col, value=value)

    _autosize(ws, [26, 24, 11, 18, 9, 15, 18, 11, 18, 20, 22])
    ws.freeze_panes = ws.cell(row=4, column=1)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
